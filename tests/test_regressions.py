"""Regression checks use a disposable database and synthetic customers only."""
import concurrent.futures
import io
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from contextlib import closing
from pypdf import PdfReader

from werkzeug.security import generate_password_hash

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ['FT_SECRET_KEY'] = 'test-only-secret-key'
os.environ['FT_LOGIN_USERNAME'] = 'test-admin'
os.environ['FT_LOGIN_PASSWORD_HASH'] = generate_password_hash('test-password')
os.environ['FT_ADMIN_PASSWORD_HASH'] = generate_password_hash('test-report-password')
import app as application
from billing_math import totals
from openpyxl import load_workbook


class BillingRegressionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.old_db = application.DB_PATH
        self.old_static = application.app.static_folder
        application.DB_PATH = str(Path(self.temp.name) / 'test.db')
        application.app.static_folder = str(Path(self.temp.name) / 'static')
        application.app.config['TESTING'] = True
        self.addCleanup(setattr, application, 'DB_PATH', self.old_db)
        self.addCleanup(setattr, application.app, 'static_folder', self.old_static)
        with closing(sqlite3.connect(application.DB_PATH)) as db:
            db.executescript((ROOT / 'tests/schema.sql').read_text())
            db.execute("INSERT INTO customers(id,name,phone) VALUES (1,'Synthetic Customer','9000000000')")
            db.execute("INSERT INTO services(id,name,rate) VALUES (1,'Wash KG',100),(2,'Prepaid Plan',1000)")
            db.execute('INSERT INTO subscription_products(service_id,credit_amount) VALUES (2,1200)')
            db.execute('INSERT INTO expense_types(id,name) VALUES (1,\'Supplies\')')
            db.commit()
        self.client = self.staff_client()

    def staff_client(self):
        client = application.app.test_client()
        with client.session_transaction() as session:
            session['user'] = 'test-admin'
            session['csrf_token'] = 'test-csrf'
        return client

    def post(self, url, client=None, **kwargs):
        return (client or self.client).post(url, headers={'X-CSRF-Token':'test-csrf'}, **kwargs)

    def sql(self, query, args=()):
        with closing(sqlite3.connect(application.DB_PATH)) as db:
            db.row_factory = sqlite3.Row
            result = db.execute(query, args).fetchall()
            db.commit()
        return result

    def wallet(self, amount):
        self.sql('INSERT OR REPLACE INTO customer_subscriptions(customer_id,current_balance,is_active) VALUES (1,?,1)', (amount,))

    def balance(self):
        return self.sql('SELECT current_balance FROM customer_subscriptions WHERE customer_id=1')[0][0]

    def bill(self, service='Wash KG', paid=False, **overrides):
        data = dict(customerPhone='9000000000',customerName='Synthetic Customer',pickupAddress='Test address',
                    orderType='Walk-in',deliveryDate='2026-10-26',discountType='Rs',discountValue='0',
                    surchargeType='Rs',surchargeValue='0',advancePaid='0',
                    **{'service[]':[service], 'qty[]':['1'], 'rate[]':['500' if service=='Wash KG' else '1000'], 'total[]':['1'], 'pieces[]':['3']})
        if paid:
            data['isPaid']='on'
        data.update(overrides)
        response = self.post('/billing', data=data)
        self.assertEqual(response.status_code, 302, response.get_data(as_text=True))
        return int(response.location.rsplit('/',1)[1])

    def info(self, bill_id):
        token = self.sql('SELECT token FROM bills WHERE id=?',(bill_id,))[0][0]
        return self.client.get('/api/bill-info/'+token).get_json()

    def test_server_recomputes_submitted_line_total(self):
        bill_id = self.bill()
        self.assertEqual(self.sql('SELECT total FROM bills WHERE id=?',(bill_id,))[0][0],500)
        self.assertEqual(self.sql('SELECT total FROM bill_items WHERE bill_id=?',(bill_id,))[0][0],500)

    def test_amounts_agree_across_invoice_api_and_customer_page(self):
        self.wallet(200)
        bill_id = self.bill(advancePaid='100')
        info = self.info(bill_id)
        self.assertEqual((info['total_before_wallet'],info['wallet_used'],info['final_payable'],info['advance_paid'],info['amount_due']),(500,200,300,100,200))
        with application.app.test_request_context('/'):
            context = application.invoice_context(bill_id=bill_id, make_qr=False)
            self.assertEqual(context['amount_due'],200)
        page = self.client.get('/pay/'+info['bill']['token'])
        self.assertEqual(page.status_code,200)
        self.assertIn(b'am=200.00',page.data)
        self.assertIn(b'Wallet Used',page.data)
        for url in (f'/bill/{bill_id}', f'/bill/{bill_id}/embed'):
            page=self.client.get(url)
            self.assertEqual(page.status_code,200)
            self.assertIn(b'200.00',page.data)

    def test_discount_express_wallet_advance(self):
        self.wallet(200)
        bid=self.bill(discountType='%',discountValue='10',expressService='on',surchargeType='%',surchargeValue='20',advancePaid='50')
        info=self.info(bid)
        self.assertEqual(info['total_before_wallet'],550)
        self.assertEqual(info['amount_due'],300)
        self.assertEqual(self.balance(),0)

    def test_decimal_rounding(self):
        bid=self.bill(**{'qty[]':['1.25'],'rate[]':['10.10']})
        self.assertEqual(self.info(bid)['amount_due'],12.63)

    def test_wallet_never_pays_for_topup(self):
        self.wallet(250)
        bid=self.bill(service='Prepaid Plan')
        self.assertEqual(self.info(bid)['wallet_used'],0)
        self.assertEqual(self.balance(),250)

    def test_paid_at_creation_unlocks_credit_once(self):
        bid=self.bill(service='Prepaid Plan',paid=True)
        self.assertEqual(self.balance(),1200)
        for _ in range(2):
            self.assertEqual(self.post('/api/mark-paid',json={'bill_id':bid}).status_code,200)
        self.assertEqual(self.balance(),1200)
        self.assertEqual(len(self.sql("SELECT * FROM subscription_ledger WHERE bill_id=? AND txn_type='CREDIT'",(bid,))),1)
        self.assertEqual(self.info(bid)['amount_due'],0)

    def test_pending_topup_waits_for_payment(self):
        bid=self.bill(service='Prepaid Plan')
        self.assertEqual(self.sql('SELECT * FROM customer_subscriptions'),[])
        self.post('/api/mark-paid',json={'bill_id':bid})
        self.assertEqual(self.balance(),1200)

    def test_cancel_wallet_bill_refunds_once(self):
        self.wallet(600)
        bid=self.bill()
        self.assertEqual(self.balance(),100)
        for _ in range(2):
            self.assertEqual(self.post('/api/cancel-bill',json={'bill_id':bid}).status_code,200)
        self.assertEqual(self.balance(),600)
        self.assertEqual(self.info(bid)['amount_due'],0)
        page=self.client.get('/pay/'+self.info(bid)['bill']['token'])
        self.assertIn(b'CANCELLED',page.data)
        self.assertNotIn(b'upi://',page.data)

    def test_cancel_topup_removes_credit_once(self):
        self.wallet(50)
        bid=self.bill(service='Prepaid Plan',paid=True)
        self.assertEqual(self.balance(),1250)
        for _ in range(2):
            self.assertEqual(self.post('/api/cancel-bill',json={'bill_id':bid}).status_code,200)
        self.assertEqual(self.balance(),50)
        self.assertEqual(len(self.sql("SELECT * FROM subscription_ledger WHERE bill_id=? AND txn_type='ADJUST'",(bid,))),1)

    def test_cancel_spent_topup_is_atomic_and_blocked(self):
        bid=self.bill(service='Prepaid Plan',paid=True)
        self.bill()
        before=self.balance()
        response=self.post('/api/cancel-bill',json={'bill_id':bid})
        self.assertEqual(response.status_code,409)
        self.assertEqual(self.balance(),before)
        self.assertEqual(self.sql('SELECT void FROM bills WHERE id=?',(bid,))[0][0],0)
        self.assertEqual(self.sql("SELECT * FROM subscription_ledger WHERE bill_id=? AND txn_type='ADJUST'",(bid,)),[])

    def test_cancelled_bill_cannot_be_paid_or_delivered(self):
        bid=self.bill(service='Prepaid Plan')
        self.post('/api/cancel-bill',json={'bill_id':bid})
        for endpoint in ('mark-paid','mark-delivered'):
            self.assertEqual(self.post('/api/'+endpoint,json={'bill_id':bid}).status_code,409)
        self.assertEqual(self.sql('SELECT * FROM customer_subscriptions'),[])

    def test_missing_bill_returns_404(self):
        for endpoint in ('mark-paid','mark-delivered','cancel-bill'):
            self.assertEqual(self.post('/api/'+endpoint,json={'bill_id':999}).status_code,404)
        for url in ('/bill/999','/invoice/pdf/999','/api/bill-info/missing','/pay/missing'):
            self.assertEqual(self.client.get(url).status_code,404)

    def test_edit_keeps_surcharge_and_wallet(self):
        self.wallet(200)
        bid=self.bill(expressService='on',surchargeValue='50',advancePaid='50')
        response=self.post(f'/bill/edit/{bid}',data={'discountType':'Rs','discountValue':'25'})
        self.assertEqual(response.status_code,302)
        self.assertEqual(self.info(bid)['amount_due'],275)
        self.assertEqual(self.sql('SELECT balance_amount FROM bills WHERE id=?',(bid,))[0][0],275)

    def test_edit_cannot_overwrite_paid_or_cancelled_bill(self):
        for paid in (True,False):
            bid=self.bill(paid=paid)
            if not paid: self.post('/api/cancel-bill',json={'bill_id':bid})
            self.assertEqual(self.post(f'/bill/edit/{bid}',data={'discountValue':'1'}).status_code,409)

    def test_edit_cannot_hide_wallet_already_used(self):
        self.wallet(500)
        bid=self.bill()
        self.assertEqual(self.post(f'/bill/edit/{bid}',data={'discountType':'Rs','discountValue':'100'}).status_code,400)
        self.assertEqual(self.balance(),0)

    def test_sales_are_full_bill_value_not_wallet_or_advance(self):
        self.wallet(200)
        bid=self.bill(discountValue='50',expressService='on',surchargeValue='25',advancePaid='100')
        day=self.info(bid)['bill']['date'][:10]
        for url in (f'/api/summary-daily?date={day}',f'/api/summary-range?start={day}&end={day}'):
            self.assertEqual(self.client.get(url).get_json()['total_sales'],'475.00')
        summary=self.client.get(f'/api/range-summary-by-date?start={day}&end={day}').get_json()[0]
        self.assertEqual(summary['walkin_sales'],475)
        self.assertEqual(self.client.get(f'/api/range-report?start={day}&end={day}').get_json()[0]['total'],475)
        self.post('/api/cancel-bill',json={'bill_id':bid})
        self.assertEqual(self.client.get(f'/api/summary-daily?date={day}').get_json()['total_sales'],'0.00')

    def test_staff_api_routes_are_not_public(self):
        client=application.app.test_client()
        for rule in application.app.url_map.iter_rules():
            if rule.rule.startswith('/api/') and rule.endpoint != 'api_bill_info':
                url=rule.rule.replace('<int:type_id>','1').replace('<int:expense_id>','1').replace('<int:bill_id>','1')
                method=next(m for m in ('GET','POST','PUT','DELETE') if m in rule.methods)
                response=client.open(url,method=method,json={} if method!='GET' else None)
                self.assertEqual(response.status_code,401,(url,response.status_code))

    def test_public_customer_link_still_works_without_login(self):
        bid=self.bill()
        token=self.info(bid)['bill']['token']
        client=application.app.test_client()
        self.assertEqual(client.get('/pay/'+token).status_code,200)
        self.assertEqual(client.get('/api/bill-info/'+token).status_code,200)
        self.assertEqual(client.get('/api/bill-info/'+token).headers['Cache-Control'],'no-store')

    def test_admin_password_is_enforced_on_server(self):
        self.assertEqual(self.client.get('/api/outstanding-report').status_code,403)
        self.assertEqual(self.post('/api/verify-password',json={'password':'wrong'}).status_code,403)
        self.assertEqual(self.post('/api/verify-password',json={'password':'test-report-password'}).status_code,200)
        self.assertEqual(self.client.get('/api/outstanding-report').status_code,200)

    def test_csrf_and_cross_origin_changes_blocked(self):
        self.assertEqual(self.client.post('/api/cancel-bill',json={'bill_id':1}).status_code,403)
        self.assertEqual(self.client.post('/api/cancel-bill',json={'bill_id':1},headers={'X-CSRF-Token':'test-csrf','Origin':'https://other.example'}).status_code,403)
        self.assertEqual(self.client.get('/services?delete=1').status_code,405)
        self.assertEqual(self.client.get('/customers?delete=1').status_code,405)

    def test_login_uses_configured_hash_and_issues_new_session(self):
        client=application.app.test_client()
        page=client.get('/login')
        self.assertEqual(page.status_code,200)
        self.assertIn(b'request_security.js',page.data)
        with client.session_transaction() as session: token=session['csrf_token']
        response=client.post('/login',data={'username':'test-admin','password':'test-password','csrf_token':token})
        self.assertEqual(response.status_code,302)
        self.assertEqual(client.get('/').status_code,200)

    def test_bad_amounts_never_create_partial_bills(self):
        for overrides in ({'qty[]':['-1']},{'rate[]':['NaN']},{'advancePaid':'-5'},{'qty[]':['0']},{'pieces[]':['2.5']},{'discountType':'%','discountValue':'101'},{'service[]':[]},{'rate[]':[]}):
            before=len(self.sql('SELECT * FROM bills'))
            data={'customerPhone':'9000000000','customerName':'Test','pickupAddress':'','service[]':['Wash KG'],'qty[]':['1'],'rate[]':['10']}
            data.update(overrides)
            self.assertEqual(self.post('/billing',data=data).status_code,400,overrides)
            self.assertEqual(len(self.sql('SELECT * FROM bills')),before)

    def test_failed_write_rolls_back_customer_and_bill(self):
        with closing(sqlite3.connect(application.DB_PATH)) as db:
            db.execute("CREATE TRIGGER reject_item BEFORE INSERT ON bill_items BEGIN SELECT RAISE(ABORT,'test failure'); END")
            db.commit()
        response=self.post('/billing',data={'customerPhone':'9111111111','customerName':'Not saved','service[]':['Wash KG'],'qty[]':['1'],'rate[]':['100']})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.sql('SELECT * FROM bills'),[])
        self.assertEqual(self.sql("SELECT * FROM customers WHERE phone='9111111111'"),[])

    def test_exports_have_headers_and_exclude_cancelled_bills(self):
        bid=self.bill()
        day=self.info(bid)['bill']['date'][:10]
        self.post('/api/cancel-bill',json={'bill_id':bid})
        for url in (f'/reports/export-daily?date={day}',f'/reports/export?start={day}&end={day}'):
            response=self.client.get(url)
            self.assertEqual(response.status_code,200)
            workbook=load_workbook(io.BytesIO(response.data))
            self.assertEqual(workbook.active.cell(1,1).value,'bill_id')
            self.assertEqual(workbook.active.max_row,1)

    def test_pdf_download_is_real_and_matches_shared_amounts(self):
        self.wallet(200)
        bid=self.bill(advancePaid='100')
        response=self.client.get(f'/invoice/pdf/{bid}')
        self.assertEqual(response.status_code,200)
        self.assertTrue(response.data.startswith(b'%PDF-'))
        reader=PdfReader(io.BytesIO(response.data))
        text=''.join(page.extract_text() for page in reader.pages)
        self.assertIn('Amount Due',text)
        self.assertIn('INR 200.00',text)
        self.assertIn('Synthetic Customer',text)
        self.assertIn('Wallet Used',text)

    def test_pdf_supports_many_items_and_escapes_customer_text(self):
        bid=self.bill()
        self.sql("UPDATE customers SET name='<Customer & Test>' WHERE id=1")
        for _ in range(70):
            self.sql("INSERT INTO bill_items(bill_id,service_id,service_name,qty,rate,total,items) VALUES (?,1,'Long laundry service description',1,1,1,1)",(bid,))
        response=self.client.get(f'/invoice/pdf/{bid}')
        reader=PdfReader(io.BytesIO(response.data))
        self.assertGreater(len(reader.pages),1)
        self.assertIn('<Customer & Test>',reader.pages[0].extract_text())

    def test_edit_fragment_has_csrf_field_without_javascript(self):
        bid=self.bill()
        response=self.client.get(f'/bill/edit/{bid}')
        self.assertIn(b'name="csrf_token"',response.data)

    def test_paid_delivery_outstanding_row_has_zero_due(self):
        bid=self.bill()
        self.post('/api/mark-paid',json={'bill_id':bid})
        self.post('/api/verify-password',json={'password':'test-report-password'})
        data=self.client.get('/api/outstanding-report').get_json()
        self.assertEqual(data[0]['balance_amount'],0)
        self.post('/api/mark-delivered',json={'bill_id':bid})
        self.assertEqual(self.client.get('/api/outstanding-report').get_json(),[])

    def test_concurrent_payment_calls_credit_once(self):
        bid=self.bill(service='Prepaid Plan')
        def pay(_):
            return self.post('/api/mark-paid',client=self.staff_client(),json={'bill_id':bid}).status_code
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            statuses=list(pool.map(pay,range(2)))
        self.assertEqual(statuses,[200,200])
        self.assertEqual(self.balance(),1200)

    def test_concurrent_bill_wallet_debits_cannot_overspend(self):
        self.wallet(600)
        def create(_):
            return self.post('/billing',client=self.staff_client(),data={'customerPhone':'9000000000','customerName':'Synthetic Customer','service[]':['Wash KG'],'qty[]':['1'],'rate[]':['500']}).status_code
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(list(pool.map(create,range(2))),[302,302])
        self.assertEqual(self.balance(),0)
        self.assertEqual(self.sql("SELECT SUM(amount) FROM subscription_ledger WHERE txn_type='DEBIT'")[0][0],600)
        self.assertEqual(self.sql('SELECT SUM(balance_amount) FROM bills')[0][0],400)

    def test_concurrent_cancellation_refunds_once(self):
        self.wallet(600)
        bid=self.bill()
        def cancel(_):
            return self.post('/api/cancel-bill',client=self.staff_client(),json={'bill_id':bid}).status_code
        with concurrent.futures.ThreadPoolExecutor(max_workers=2) as pool:
            self.assertEqual(list(pool.map(cancel,range(2))),[200,200])
        self.assertEqual(self.balance(),600)

    def test_expense_validation_and_invalid_type(self):
        response=self.post('/api/expenses',json={'expense_type_id':1,'amount':'-1','date':'2026-10-25'})
        self.assertEqual(response.status_code,400)
        response=self.post('/api/expenses',json={'expense_type_id':999,'amount':'1','date':'2026-10-25'})
        self.assertEqual(response.status_code,400)
        self.assertEqual(self.sql('SELECT * FROM expenses'),[])
        response=self.post('/api/expenses',json={'expense_type_id':1,'amount':'25.25','date':'2026-10-25'})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.sql('SELECT amount FROM expenses')[0][0],25.25)

    def test_static_financial_exports_are_not_public(self):
        client=application.app.test_client()
        self.assertEqual(client.get('/static/reports/old.xlsx').status_code,401)
        self.assertEqual(client.get('/static/invoices/old.pdf').status_code,401)

    def test_amount_discount_remains_selected_when_editing(self):
        bid=self.bill(discountType='Rs',discountValue='25')
        self.assertIn(b'value="Amount" selected',self.client.get(f'/bill/edit/{bid}').data)

    def test_wallet_float_residue_cannot_break_a_full_payment(self):
        self.wallet(0.29999999999999993)
        bid=self.bill(**{'rate[]':['0.30']})
        self.assertEqual(self.balance(),0)
        self.assertEqual(self.info(bid)['amount_due'],0)

    def test_public_https_origin_works_behind_http_reverse_proxy(self):
        old = application.app.config.get('PUBLIC_ORIGIN')
        application.app.config['PUBLIC_ORIGIN'] = 'https://staff.example.test'
        try:
            response = self.client.post('/api/verify-password', json={'password':'test-report-password'},
                headers={'X-CSRF-Token':'test-csrf','Origin':'https://staff.example.test'})
            self.assertEqual(response.status_code,200)
            response = self.client.post('/api/verify-password', json={'password':'test-report-password'},
                headers={'X-CSRF-Token':'test-csrf','Origin':'https://attacker.example.test'})
            self.assertEqual(response.status_code,403)
        finally:
            if old is None: application.app.config.pop('PUBLIC_ORIGIN',None)
            else: application.app.config['PUBLIC_ORIGIN']=old

    def test_fake_generation_endpoint_does_not_claim_saved_bill(self):
        response=self.post('/api/generate-bill',json={})
        self.assertEqual(response.status_code,501)
        self.assertFalse(response.get_json()['success'])
        self.assertEqual(self.sql('SELECT * FROM bills'),[])

    def test_existing_stored_line_totals_are_preserved(self):
        result=totals({'discount_type':'Rs','discount_value':0},[{'qty':2,'rate':100,'total':150}])
        self.assertEqual(result['subtotal'],150)


if __name__=='__main__':
    unittest.main()
